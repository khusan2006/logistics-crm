"""A bron of several interchangeable markalar, one kg shared by all of them.

Some granula stand in for each other, and a mijoz booking 20 000 kg of "A or B" takes
whichever lands. The markalar are equal, each has its own narx, and a sotuv of any of
them draws the same kg down.
"""
from decimal import Decimal

from conftest import bron_data, line_data, make_bron
from crm.models import (Currency, CustomerPayment, Reservation, Sale, bron_advance_holds,
                        bron_queue, brand_reserved_kg, brand_stock_costed)
from tests.test_reservations import _arrived_lot, _arrived_lot_for, _customer, _in_transit_lot


def _reserve(client, customer, brands, prices, kg="20000", currency="usd"):
    return client.post("/reservations/new/", bron_data({
        "customer": customer.pk, "brand": brands, "price": prices, "kg": kg,
        "currency": currency, "exchange_rate": "12000", "note": ""}))


def _sell(client, customer, brand, kg, price="1.50"):
    return client.post("/sales/new/", {
        "customer": customer.pk, "currency": "usd", "exchange_rate": "12000",
        "date": "2026-07-20",
        **line_data({"brand": brand, "kg": kg, "price": price})})


class TestTakingOne:
    def test_one_bron_names_both_markalar_at_their_own_narx(self, admin_client, db):
        _arrived_lot(kg="10000", brand="LLDPE")
        _in_transit_lot(kg="5000", brand="HDPE")
        resp = _reserve(admin_client, _customer(), ["LLDPE", "HDPE"], ["1.40", "1.55"])
        assert resp.status_code == 302
        bron = Reservation.objects.get()
        assert bron.kg == Decimal("20000.000")
        assert [(i.brand, i.price) for i in bron.item_list] == [
            ("LLDPE", Decimal("1.4000")), ("HDPE", Decimal("1.5500"))]
        assert bron.items.get(brand="HDPE").price_uzs == Decimal("18600.00")

    def test_a_marka_may_be_left_without_a_narx(self, admin_client, db):
        _arrived_lot(kg="10000", brand="LLDPE")
        _in_transit_lot(kg="5000", brand="HDPE")
        _reserve(admin_client, _customer(), ["LLDPE", "HDPE"], ["1.40", ""])
        bron = Reservation.objects.get()
        assert bron.items.get(brand="HDPE").price is None
        # Which marka the mijoz takes is not known, so neither is the value.
        assert bron.value_price is None and bron.totals == []

    def test_the_same_marka_twice_is_refused(self, admin_client, db):
        _arrived_lot(kg="10000", brand="LLDPE")
        resp = _reserve(admin_client, _customer(), ["LLDPE", "LLDPE"], ["1.40", "1.50"])
        assert resp.status_code == 200
        assert "Bu marka ro&#x27;yxatda bor" in resp.content.decode()
        assert not Reservation.objects.exists()

    def test_no_marka_at_all_is_refused(self, admin_client, db):
        resp = _reserve(admin_client, _customer(), [], [])
        assert resp.status_code == 200
        assert "Kamida bitta marka" in resp.content.decode()
        assert not Reservation.objects.exists()

    def test_a_narx_with_no_kurs_is_refused(self, admin_client, db):
        _arrived_lot(kg="10000", brand="LLDPE")
        resp = admin_client.post("/reservations/new/", bron_data({
            "customer": _customer().pk, "brand": ["LLDPE"], "price": ["1.40"],
            "kg": "1000", "currency": "usd", "exchange_rate": "", "note": ""}))
        assert resp.status_code == 200
        assert "Dollar kursini kiriting" in resp.content.decode()
        assert not Reservation.objects.exists()

    def test_a_som_bron_keeps_each_typed_som_narx(self, admin_client, db):
        _arrived_lot(kg="10000", brand="LLDPE")
        _in_transit_lot(kg="5000", brand="HDPE")
        _reserve(admin_client, _customer(), ["LLDPE", "HDPE"], ["16800", "18000"],
                 currency="uzs")
        bron = Reservation.objects.get()
        assert [i.price_own for i in bron.item_list] == [Decimal("16800.00"),
                                                        Decimal("18000.00")]


class TestEditingIt:
    def test_a_marka_can_be_added_and_the_edit_form_shows_the_som_figure(
            self, admin_client, db):
        _arrived_lot(kg="10000", brand="LLDPE")
        _in_transit_lot(kg="5000", brand="HDPE")
        customer = _customer()
        _reserve(admin_client, customer, ["LLDPE"], ["16800"], currency="uzs")
        bron = Reservation.objects.get()
        html = admin_client.get(f"/reservations/{bron.pk}/edit/").content.decode()
        assert 'name="items-0-price"' in html and "16800" in html

        resp = admin_client.post(f"/reservations/{bron.pk}/edit/", bron_data({
            "customer": customer.pk, "brand": ["LLDPE", "HDPE"],
            "price": ["16800", "18000"], "kg": "20000", "currency": "uzs",
            "exchange_rate": "12000", "note": ""}, bron))
        assert resp.status_code == 302
        assert bron.brands == ["LLDPE", "HDPE"]
        assert bron.items.get(brand="LLDPE").price_uzs == Decimal("16800.00")

    def test_a_marka_can_be_struck_off(self, admin_client, db):
        _arrived_lot(kg="10000", brand="LLDPE")
        _in_transit_lot(kg="5000", brand="HDPE")
        customer = _customer()
        _reserve(admin_client, customer, ["LLDPE", "HDPE"], ["1.40", "1.55"])
        bron = Reservation.objects.get()
        resp = admin_client.post(f"/reservations/{bron.pk}/edit/", bron_data({
            "customer": customer.pk, "brand": ["LLDPE"], "price": ["1.40"],
            "kg": "20000", "currency": "usd", "exchange_rate": "12000", "note": ""},
            bron))
        assert resp.status_code == 302
        assert bron.brands == ["LLDPE"]


class TestServingIt:
    def test_any_marka_draws_the_one_kg(self, admin_client, db):
        """12 000 of one and 8 000 of the other serve a 20 000 kg bron in full."""
        a = _arrived_lot_for("Pars", kg="15000", brand="LLDPE")
        b = _arrived_lot_for("Tabriz", kg="15000", brand="HDPE")
        customer = _customer()
        bron = make_bron(customer, ["LLDPE", "HDPE"], kg="20000", price=["1.40", "1.55"])

        _sell(admin_client, customer, b.brand, "12000")
        bron.refresh_from_db()
        assert bron.remaining_kg == Decimal("8000.000")
        assert bron.status == Reservation.Status.ACTIVE

        _sell(admin_client, customer, a.brand, "8000")
        bron.refresh_from_db()
        assert bron.remaining_kg == Decimal("0.000")
        assert bron.status == Reservation.Status.CONVERTED
        assert {s.reservation_id for s in Sale.objects.all()} == {bron.pk}

    def test_a_marka_it_does_not_name_leaves_it_alone(self, admin_client, db):
        _arrived_lot_for("Pars", kg="15000", brand="LLDPE")
        other = _arrived_lot_for("Tabriz", kg="15000", brand="PP")
        customer = _customer()
        bron = make_bron(customer, ["LLDPE", "HDPE"], kg="20000")
        _sell(admin_client, customer, other.brand, "5000")
        bron.refresh_from_db()
        assert bron.fulfilled_kg == Decimal("0.000")

    def test_the_sotuv_form_asks_brondan_ushlansin_for_either_marka(self, admin_client, db):
        _arrived_lot(kg="10000", brand="LLDPE")
        customer = _customer()
        make_bron(customer, ["LLDPE", "HDPE"], kg="20000")
        html = admin_client.get("/sales/new/").content.decode()
        assert "data-bron-brands=\"[&quot;HDPE&quot;, &quot;LLDPE&quot;]\"" in html

    def test_an_earlier_sotuv_of_either_marka_can_be_counted_in(self, admin_client, db):
        lot = _arrived_lot(kg="10000", brand="HDPE")
        customer = _customer()
        Sale.objects.create(customer=customer, line=lot, kg=Decimal("3000"),
                            price=Decimal("1.5"), date="2026-07-10")
        bron = make_bron(customer, ["LLDPE", "HDPE"], kg="20000")
        rows = admin_client.get("/reservations/").context["rows"]
        assert rows[0].countable_sales
        resp = admin_client.post(f"/reservations/{bron.pk}/count-sales/",
                                 {"sales": [Sale.objects.get().pk]})
        assert resp.status_code in (204, 302)
        bron.refresh_from_db()
        assert bron.remaining_kg == Decimal("17000.000")


class TestWhatTheScreensSay:
    def test_it_stands_in_the_queue_of_each_marka(self, db):
        customer = _customer()
        bron = make_bron(customer, ["LLDPE", "HDPE"], kg="20000")
        assert bron_queue("LLDPE") == [bron]
        assert bron_queue("HDPE") == [bron]

    def test_the_ombor_counts_it_under_both(self, db):
        _arrived_lot_for("Pars", kg="15000", brand="LLDPE")
        _arrived_lot_for("Tabriz", kg="15000", brand="HDPE")
        make_bron(_customer(), ["LLDPE", "HDPE"], kg="20000")
        assert brand_reserved_kg("LLDPE") == Decimal("20000")
        assert brand_reserved_kg("HDPE") == Decimal("20000")
        reserved = {row["brand"]: row["reserved"] for row in brand_stock_costed()}
        assert reserved == {"LLDPE": Decimal("20000"), "HDPE": Decimal("20000")}

    def test_the_ombor_row_says_the_kg_is_shared(self, admin_client, db):
        _arrived_lot_for("Pars", kg="15000", brand="LLDPE")
        make_bron(_customer(), ["LLDPE", "HDPE"], kg="20000", price=["1.40", "1.55"])
        html = admin_client.get("/ombor/").content.decode()
        assert "umumiy: LLDPE / HDPE" in html

    def test_the_list_shows_each_marka_its_narx_and_its_place(self, admin_client, db):
        _arrived_lot_for("Pars", kg="15000", brand="LLDPE")
        first = make_bron(_customer("Birinchi"), ["HDPE"], kg="1000")
        second = make_bron(_customer("Ikkinchi"), ["LLDPE", "HDPE"], kg="20000",
                           price=["1.40", "1.55"])
        rows = {r.pk: r for r in admin_client.get("/reservations/").context["rows"]}
        places = [(m["item"].brand, m["queue_pos"]) for m in rows[second.pk].marka_rows]
        assert places == [("LLDPE", 1), ("HDPE", 2)]
        assert rows[second.pk].queue_pos == 1
        # The shortcut opens on the marka actually on the shelf.
        assert rows[second.pk].serve_item.brand == "LLDPE"
        assert rows[first.pk].queue_pos == 1
        html = admin_client.get("/reservations/").content.decode()
        assert "2 marka uchun umumiy" in html

    def test_jami_is_the_range_its_narxlar_come_to(self, db):
        bron = make_bron(_customer(), ["LLDPE", "HDPE"], kg="1000",
                         price=["1.40", "1.55"])
        assert [usd for usd, _uzs in bron.totals] == [Decimal("1400.00"),
                                                      Decimal("1550.00")]

    def test_the_avans_is_held_at_the_dearest_narx(self, db):
        customer = _customer()
        CustomerPayment.objects.create(
            customer=customer, date="2026-07-01", amount=Decimal("10000"),
            amount_uzs=Decimal("120000000"), currency=Currency.USD, method="cash")
        bron = make_bron(customer, ["LLDPE", "HDPE"], kg="1000", price=["1.40", "1.55"],
                         status=Reservation.Status.ACTIVE)
        hold = bron_advance_holds([bron])[bron.pk]
        assert hold["held"] == Decimal("1550.00")
        assert hold["short"] == Decimal("0.00")

    def test_the_excel_file_spells_out_differing_narxlar(self, admin_client, db):
        from io import BytesIO

        import openpyxl

        make_bron(_customer(), ["LLDPE", "HDPE"], kg="1000", price=["1.40", "1.55"])
        resp = admin_client.get("/reservations/export.xlsx")
        ws = openpyxl.load_workbook(BytesIO(resp.content)).worksheets[0]
        headers = [c.value for c in ws[1]]
        row = next(ws.iter_rows(min_row=2, values_only=True))
        assert row[headers.index("Marka")] == "LLDPE / HDPE"
        assert row[headers.index("Narx ($)")] is None
        assert row[headers.index("Markalar narxi")] == "LLDPE: 1.4; HDPE: 1.55"


class TestMergeBrand:
    def test_a_bron_naming_both_spellings_keeps_one(self, admin_client, db):
        from io import StringIO

        from django.core.management import call_command

        _arrived_lot_for("Pars", kg="15000", brand="i 1561")
        _arrived_lot_for("Tabriz", kg="15000", brand="и 1561")
        bron = make_bron(_customer(), ["и 1561", "i 1561"], kg="20000")
        call_command("merge_brand", "и 1561", into="i 1561", apply=True, stdout=StringIO())
        assert bron.brands == ["i 1561"]
