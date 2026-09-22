"""Birja — granula bought on the exchange here, rather than agreed with a hamkor in
Eron.

The whole feature rests on one decision: a birja kelishuv is an ordinary `Contract`
and a birja yuk an ordinary `Shipment`, both hanging off a singleton hamkor row
flagged `is_birja`. That is what lets the ombor, the sotuvlar and the kassa stay one
set of books — and it is what these tests are mostly checking: that the two lists
separate cleanly in BOTH directions, and that the Eron-road facts a birja load has
no business carrying (a QR kod, a bojxona) stay off it.
"""
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from conftest import line_data, make_shipment, supplier_payment_rows

from crm.models import (
    Contract, ContractExpense, ContractLine, Currency, Partner, Shipment,
    ShipmentExpense, ShipmentStatus, birja_partner, brand_stock_costed,
)

pytestmark = pytest.mark.django_db


# --- helpers ---------------------------------------------------------------

def _birja_contract(brand="LLDPE", kg="1000", price="1.00", **kw):
    contract = Contract.objects.create(partner=birja_partner(),
                                       created=kw.pop("created", "2026-07-01"), **kw)
    ContractLine.objects.create(contract=contract, brand=brand,
                                kg=Decimal(kg), price=Decimal(price))
    return contract


def _hamkor_contract(brand="HDPE", kg="1000", price="1.00"):
    partner = Partner.objects.create(name="Pars", phone="1", city="Tehron")
    contract = Contract.objects.create(partner=partner, created="2026-07-01")
    ContractLine.objects.create(contract=contract, brand=brand,
                                kg=Decimal(kg), price=Decimal(price))
    return contract


def _first_birja_status():
    return ShipmentStatus.for_kind(birja=True).first()


def _post_birja_shipment(client, contract, **extra):
    row = {"contract_line": contract.lines.first().pk, "kg": extra.pop("kg", "400")}
    data = {"contract": contract.pk, "status": _first_birja_status().pk,
            "sent": "2026-07-05", "eta": "2026-07-20", "transport": "01A111AA",
            "note": "", **line_data(row)}
    data.update(extra)
    return client.post("/birja/yuklar/new/", data)


# --- the counterparty and the kod ------------------------------------------

def test_the_birja_hamkor_is_one_row_created_on_first_use():
    """Lazily, and never twice. It is not seeded by a migration on purpose:
    `wipe_business_data()` deletes every Partner, so a seeded row would vanish the
    first time starting data is reloaded and the numbering would restart at
    birja-1 beside codes that already exist."""
    assert not Partner.objects.filter(is_birja=True).exists()
    first = birja_partner()
    assert birja_partner().pk == first.pk
    assert Partner.objects.filter(is_birja=True).count() == 1


def test_kelishuvlar_are_coded_birja_1_birja_2():
    """The whole reason the counterparty is a Partner: `Contract.save()` already
    slugifies the hamkor's name and counts on `Partner.code_counter`, so the birja
    codes fall out of the existing machinery rather than a second number line."""
    assert _birja_contract(brand="A").code == "birja-1"
    assert _birja_contract(brand="B").code == "birja-2"


def test_a_deleted_birja_kelishuv_does_not_hand_its_number_back():
    """Same no-recycle rule a hamkor's codes follow — a kod that has been read out
    loud must never come back meaning something else."""
    _birja_contract(brand="A")
    _birja_contract(brand="B").delete()
    assert _birja_contract(brand="C").code == "birja-3"


def test_the_birja_hamkor_is_not_offered_as_one():
    """It is the exchange standing in as a counterparty, not somebody to strike a
    deal with — so it is off the Hamkorlar list and off the kelishuv form's picker."""
    from crm.forms import ContractForm
    birja = birja_partner()
    assert birja not in ContractForm().fields["partner"].queryset


def test_the_hamkorlar_list_does_not_show_it(admin_client):
    _birja_contract()
    Partner.objects.create(name="Pars", phone="1", city="Tehron")
    names = [p.name for p in admin_client.get("/partners/").context["page"]]
    assert names == ["Pars"]


# --- the two lists separate, both ways -------------------------------------

def test_the_two_kelishuvlar_lists_hold_only_their_own(admin_client):
    birja, hamkor = _birja_contract(), _hamkor_contract()
    assert [c.pk for c in admin_client.get("/birja/kelishuvlar/").context["rows"]] \
        == [birja.pk]
    assert [c.pk for c in admin_client.get("/contracts/").context["rows"]] \
        == [hamkor.pk]


def test_the_two_yuklar_lists_hold_only_their_own(admin_client):
    birja = make_shipment(contract=_birja_contract(), status=_first_birja_status())
    hamkor = make_shipment(contract=_hamkor_contract())
    assert [s.pk for s in admin_client.get("/birja/yuklar/").context["shipments"]] \
        == [birja.pk]
    assert [s.pk for s in admin_client.get("/shipments/").context["shipments"]] \
        == [hamkor.pk]


def _sheet_text(response):
    ws = openpyxl.load_workbook(BytesIO(response.content)).active
    return {str(cell.value) for row in ws.iter_rows() for cell in row}


def test_the_excel_button_exports_the_list_it_was_pressed_on(admin_client):
    """The flag is threaded through `_filter_contracts` / `_filter_shipments` rather
    than read off the path, so the file cannot end up holding the other list."""
    _birja_contract(brand="BIRJAMARKA")
    _hamkor_contract(brand="ERONMARKA")
    birja_file = _sheet_text(admin_client.get("/birja/kelishuvlar/export.xlsx"))
    hamkor_file = _sheet_text(admin_client.get("/contracts/export.xlsx"))
    assert "BIRJAMARKA" in birja_file and "ERONMARKA" not in birja_file
    assert "ERONMARKA" in hamkor_file and "BIRJAMARKA" not in hamkor_file


def test_the_birja_list_opens_oldest_first(admin_client):
    """The other way up from the Eron list: the oldest open birja kelishuv is the one
    whose trucks are arriving now, so it leads. Created out of date order, so the
    list cannot pass by reading the pk."""
    middle = _birja_contract(brand="B", created="2026-09-10")
    newest = _birja_contract(brand="C", created="2026-09-18")
    oldest = _birja_contract(brand="A", created="2026-09-08")
    resp = admin_client.get("/birja/kelishuvlar/")
    assert [c.pk for c in resp.context["rows"]] == [oldest.pk, middle.pk, newest.pk]
    # It is the list's normal state, so Saralash draws no chip over it.
    assert resp.context["filters"]["chips"] == []


def test_a_chosen_sort_still_wins_on_the_birja_list(admin_client):
    old = _birja_contract(brand="A", created="2026-09-08")
    new = _birja_contract(brand="B", created="2026-09-18")
    resp = admin_client.get("/birja/kelishuvlar/", {"sort": "-created"})
    assert [c.pk for c in resp.context["rows"]] == [new.pk, old.pk]
    assert [chip["label"] for chip in resp.context["filters"]["chips"]] == ["Saralash"]


def test_the_birja_excel_comes_out_in_the_page_s_order(admin_client):
    new = _birja_contract(brand="B", created="2026-09-18")
    old = _birja_contract(brand="A", created="2026-09-08")
    ws = openpyxl.load_workbook(BytesIO(
        admin_client.get("/birja/kelishuvlar/export.xlsx").content)).active
    codes = [cell.value for row in ws.iter_rows() for cell in row
             if cell.value in (old.code, new.code)]
    assert codes == [old.code, new.code]


# --- valyuta ---------------------------------------------------------------
#
# Birja purchases are struck in so'm, which is a DEFAULT and not a rule. Everything
# below is the app's existing dual-currency machinery — `own_side`, `convert_pair`,
# the locked picker — asked of a birja kelishuv, because a screen that quietly
# assumed so'm would be wrong the first time a lot is quoted in dollars.

def test_a_birja_kelishuv_can_still_be_struck_in_dollars(admin_client):
    """The so'm is where the picker STARTS, not where it is stuck."""
    resp = admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "usd", "created": "2026-07-04", "note": "",
        **line_data({"brand": "7000F", "kg": "1000", "price": "1.05"})})
    assert resp.status_code == 302
    contract = Contract.objects.get()
    assert contract.currency == Currency.USD and not contract.is_som
    assert contract.lines.first().price == Decimal("1.0500")


# ── One marka, several lots ──────────────────────────────────────────────────────
#
# The birja sells the same granula through the day at whatever it is asking, so one
# purchase is routinely "и 1561" three times at three prices. The Mahsulotlar
# formset used to key on the marka alone and rejected everything after the first
# row, leaving the operator to split one purchase across three kelishuvlar.
#
# BIRJA ONLY. A hamkor kelishuv is one agreement negotiated at one price per marka,
# so a marka repeated there is still a slip — see the last test in this block, and
# `BaseContractLineFormSet` for why the views have to hand the side down.

def test_one_marka_can_be_bought_in_lots_at_different_prices(admin_client):
    """The shape off the exchange floor: 30 000 at 16 600, 30 000 at 16 500,
    7 000 at 16 400 — one kelishuv, three lots of one granula."""
    resp = admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "и 1561", "kg": "30000", "price": "16600",
                     "planned_trucks": "2"},
                    {"brand": "и 1561", "kg": "30000", "price": "16500",
                     "planned_trucks": "2"},
                    {"brand": "и 1561", "kg": "7000", "price": "16400",
                     "planned_trucks": "1"})})
    assert resp.status_code == 302
    contract = Contract.objects.get()
    assert [ln.price_uzs for ln in contract.lines.all()] == [
        Decimal("16600.00"), Decimal("16500.00"), Decimal("16400.00")]
    # Each lot keeps its own qolgan kg — a truck is booked against the lot it was
    # priced out of, not against a marka-wide pool.
    assert [ln.remaining_kg for ln in contract.lines.all()] == [
        Decimal("30000.000"), Decimal("30000.000"), Decimal("7000.000")]


def test_the_same_marka_at_the_same_narx_is_still_a_double_entry(admin_client):
    """The narx is what makes two rows two lots. Without one they are the same row
    typed twice, which is the mistake the rule was always for."""
    resp = admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "и 1561", "kg": "30000", "price": "16600"},
                    {"brand": "И 1561 ", "kg": "7000", "price": "16600"})})
    assert resp.status_code == 200
    assert not Contract.objects.exists()
    assert "shu narxda" in resp.content.decode()


def test_the_same_narx_typed_two_ways_is_still_the_same_narx(admin_client):
    """16600 and 16600.0 are one price, so they are one row twice.

    The comparison is on what `convert_pair` STORED, not on the two strings — both
    land on the same quantized Decimal, and a rule that read the raw text would let
    a trailing zero split one lot into two."""
    resp = admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "и 1561", "kg": "30000", "price": "16600"},
                    {"brand": "и 1561", "kg": "7000", "price": "16600.0"})})
    assert resp.status_code == 200
    assert not Contract.objects.exists()


def test_a_marka_taken_in_lots_is_named_once_in_the_summary(admin_client, db):
    """`brand_summary` heads dropdowns, the yuklar list and every audit line.
    "и 1561, и 1561, и 1561" names one granula three times and says nothing the
    first mention did not."""
    admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "и 1561", "kg": "30000", "price": "16600"},
                    {"brand": "и 1561", "kg": "7000", "price": "16400"},
                    {"brand": "ftor oq", "kg": "1000", "price": "15000"})})
    contract = Contract.objects.get()
    # "ftor oq" was typed in lotin and is stored the way the ombor reads it.
    assert contract.brand_summary == "и 1561, фтор оқ"


def test_a_truck_is_booked_against_the_lot_it_was_priced_out_of(admin_client, db):
    """The point of keeping the lots apart: a mashina taken off the 16 500 lot
    leaves the 16 600 one untouched. One marka-wide qolgan counter could not say
    which of the two still owes a truck."""
    admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "и 1561", "kg": "30000", "price": "16600"},
                    {"brand": "и 1561", "kg": "30000", "price": "16500"})})
    contract = Contract.objects.get()
    dear, cheap = contract.lines.all()

    make_shipment(contract_line=cheap, kg="24000")

    assert dear.remaining_kg == Decimal("30000.000")
    assert cheap.remaining_kg == Decimal("6000.000")


def test_the_ombor_still_sees_one_marka_however_many_lots_bought_it(admin_client, db):
    """Downstream nothing changes: the shelf is stocked by MARKA, and lots of one
    marka at different landed costs blend into one tannarx — the rule
    `brand_stock_costed` already followed for two kelishuvlar, now reached by two
    rows of one."""
    admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "и 1561", "kg": "30000", "price": "16600"},
                    {"brand": "и 1561", "kg": "30000", "price": "16500"})})
    contract = Contract.objects.get()
    for line in contract.lines.all():
        make_shipment(contract_line=line, kg="10000",
                      status=ShipmentStatus.arrival(), sent="2026-07-05",
                      eta="2026-07-15", arrived="2026-07-16")

    rows = brand_stock_costed()
    assert [r["brand"] for r in rows] == ["и 1561"]
    assert rows[0]["on_hand"] == Decimal("20000.000")
    # One kelishuv behind it, named once — both lots came off birja-1.
    assert rows[0]["codes"] == [contract.code]


def test_the_tolov_picker_prices_each_lot_of_a_repeated_marka(admin_client, db):
    """Which lot the money is for is a real question once one marka is on the
    kelishuv twice — and "и 1561 · 30 000 kg" twice over cannot put it."""
    from crm.forms import SupplierPaymentForm

    admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "и 1561", "kg": "30000", "price": "16600"},
                    {"brand": "и 1561", "kg": "30000", "price": "16500"})})
    field = SupplierPaymentForm().fields["contract_line"]
    labels = [field.label_from_instance(ln)
              for ln in Contract.objects.get().lines.all()]
    assert len(set(labels)) == 2
    assert all("и 1561" in label for label in labels)


def test_a_hamkor_kelishuv_still_refuses_one_marka_twice_at_any_narx(admin_client, db):
    """The rule is the birja's alone. A hamkor kelishuv is negotiated at one price
    per marka, so a second "2102 repak" is a slip however it is priced — and it
    says so in the words it always did, with no narx in them."""
    partner = Partner.objects.create(name="Pars", phone="1", city="Tehron")
    resp = admin_client.post("/contracts/new/", {
        "partner": partner.pk, "currency": "usd", "created": "2026-07-04", "note": "",
        **line_data({"brand": "2102 repak", "kg": "30000", "price": "1.05"},
                    {"brand": "2102 repak", "kg": "7000", "price": "1.02"})})
    assert resp.status_code == 200
    assert not Contract.objects.exists()
    html = resp.content.decode()
    assert "Bu mahsulot ro&#x27;yxatda bor" in html
    assert "shu narxda" not in html


def test_a_som_narx_is_stored_on_both_sides_with_the_typed_one_exact(admin_client):
    """`convert_pair` keeps the figure that was actually agreed untouched and
    computes only the other half, so a so'm narx never drifts by a tiyin — the
    dollar twin is the derivation, at the last kurs somebody really typed."""
    admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "2102 repak", "kg": "40000", "price": "13500"})})
    line = Contract.objects.get().lines.first()
    assert line.currency == Currency.UZS
    assert line.price_uzs == Decimal("13500.00")          # exactly as typed
    assert line.price == (Decimal("13500") / line.exchange_rate).quantize(
        Decimal("0.0001"))                                 # derived, not typed


def test_the_qarz_on_a_som_kelishuv_is_read_and_settled_in_som(admin_client):
    """The rule every qarz in the app follows (`own_side`): a kelishuv struck in
    so'm is owed in so'm, and a so'm to'lov closes it. Reading the dollar twin
    would leave a settled birja kelishuv permanently unfinished, because that half
    was derived at a kurs neither side agreed to."""
    contract = _birja_contract(kg="1000", price="1.00", currency="uzs")
    line = contract.lines.first()
    line.currency, line.exchange_rate = Currency.UZS, Decimal("12800")
    line.price, line.price_uzs = Decimal("1.0000"), Decimal("12800.00")
    line.save()
    make_shipment(contract=contract, kg="1000")

    assert contract.is_som
    assert contract.payable_left_uzs == Decimal("12800000.00")
    admin_client.post("/supplier-payments/new/", supplier_payment_rows(
        {"currency": "uzs", "amount": "12800000", "exchange_rate": "12800",
         "commission_percent": "", "method": "cash", "note": ""},
        contract=contract, date="2026-07-21"))
    contract = Contract.objects.get(pk=contract.pk)
    assert contract.payable_left_uzs == Decimal("0.00")
    assert contract.is_settled


def test_the_yuklar_list_leads_with_the_kelishuvs_own_currency(admin_client):
    """Qiymati is not a blended figure — the goods on a truck are priced in the one
    currency their kelishuv was struck in — so it has to lead with that side, the
    way the kelishuvlar list and every qarz already do. It led with the dollar for
    every load, which on a so'm kelishuv put a figure nobody agreed to at the front
    of the row. (A truck's XARAJAT still leads in dollars: a so'm transport bill
    beside a dollar bojxona has no agreed side.)"""
    som_contract = _birja_contract(kg="1000", price="1.00", currency="uzs")
    line = som_contract.lines.first()
    line.currency, line.exchange_rate = Currency.UZS, Decimal("12800")
    line.price, line.price_uzs = Decimal("1.0000"), Decimal("12800.00")
    line.save()
    make_shipment(contract=som_contract, kg="1000")

    html = admin_client.get("/birja/yuklar/").content.decode()
    lead = html.index("money-pair")
    alt = html.index("money-alt", lead)
    # "so'm" comes back HTML-escaped, hence the entity rather than the apostrophe.
    assert "so&#x27;m" in html[lead:alt]     # the agreed side leads
    assert "$" in html[alt:alt + 200]        # the derived twin sits beneath it


def test_a_dollar_kelishuvs_yuk_still_leads_with_the_dollar(admin_client):
    """The mirror — nothing changes for the loads that were always dollar."""
    contract = _birja_contract(kg="1000", price="1.00")   # usd by default here
    make_shipment(contract=contract, kg="1000")
    html = admin_client.get("/birja/yuklar/").content.decode()
    lead = html.index("money-pair")
    alt = html.index("money-alt", lead)
    assert "$" in html[lead:alt] and "so&#x27;m" not in html[lead:alt]


def test_the_valyuta_is_frozen_once_money_or_goods_are_on_it(admin_client):
    """Same lock a hamkor kelishuv carries: re-striking it would re-read every
    figure already booked — a 13 500 so'm narx becoming $13 500."""
    from crm.forms import ContractForm
    contract = _birja_contract(currency="uzs")
    assert not ContractForm(instance=contract, birja=True).fields["currency"].disabled
    make_shipment(contract=contract)
    assert ContractForm(instance=contract, birja=True).fields["currency"].disabled


# --- creating through the views --------------------------------------------

def test_creating_a_kelishuv_pins_the_hamkor_and_asks_nothing(admin_client):
    """No hamkor picker: there is one counterparty a birja kelishuv can have, and a
    disabled select offering it would read as a choice somebody might get wrong."""
    from crm.forms import ContractForm
    assert "partner" not in ContractForm(birja=True).fields

    resp = admin_client.post("/birja/kelishuvlar/new/", {
        "currency": "uzs", "created": "2026-07-04", "note": "",
        **line_data({"brand": "7000F", "kg": "30000", "price": "13000"})})
    assert resp.status_code == 302
    contract = Contract.objects.get()
    assert contract.partner.is_birja and contract.code == "birja-1"
    assert contract.currency == Currency.UZS


def test_a_new_kelishuv_defaults_to_som(admin_client):
    """Goods bought here are priced and settled in so'm — still a choice, not a
    lock, so a lot quoted in dollars stays recordable."""
    from crm.forms import ContractForm
    assert ContractForm(birja=True).fields["currency"].initial == Currency.UZS


def test_creating_a_yuk_routes_it_from_the_birja(admin_client):
    contract = _birja_contract()
    assert _post_birja_shipment(admin_client, contract).status_code == 302
    shipment = Shipment.objects.get()
    assert shipment.is_birja
    assert shipment.origin == "Birja"
    assert shipment.destination == "O'zbekiston"


def test_a_birja_yuk_cannot_be_loaded_against_an_eron_kelishuv(admin_client):
    """The kelishuv picker is narrowed to one side, and it is the form field that
    enforces it — posting the other side's pk is rejected, not quietly accepted."""
    hamkor = _hamkor_contract()
    resp = _post_birja_shipment(admin_client, hamkor)
    assert resp.status_code == 200 and not Shipment.objects.exists()


def test_an_eron_yuk_cannot_be_loaded_against_a_birja_kelishuv(admin_client):
    """The mirror, which is the half that is easy to leave out."""
    contract = _birja_contract()
    resp = admin_client.post("/shipments/new/", {
        "contract": contract.pk, "status": ShipmentStatus.objects.first().pk,
        "sent": "2026-07-05", "eta": "2026-07-20", "transport": "01A111AA", "note": "",
        **line_data({"contract_line": contract.lines.first().pk, "kg": "400"})})
    assert resp.status_code == 200 and not Shipment.objects.exists()


# --- the holat chains ------------------------------------------------------

def test_each_list_offers_only_its_own_chain(admin_client):
    make_shipment(contract=_birja_contract(), status=_first_birja_status())
    make_shipment(contract=_hamkor_contract())
    birja_tabs = [t["status"].name
                  for t in admin_client.get("/birja/yuklar/").context["tabs"]]
    hamkor_tabs = [t["status"].name
                   for t in admin_client.get("/shipments/").context["tabs"]]
    assert "Sotib olindi" in birja_tabs and "Chegarada" not in birja_tabs
    assert "Chegarada" in hamkor_tabs and "Sotib olindi" not in hamkor_tabs


def test_the_birja_list_opens_on_the_first_step_of_its_own_chain(admin_client):
    """There is no "Yo'lda" to open on — the operator has not named the birja steps
    yet — so it opens on whichever step is currently first."""
    make_shipment(contract=_birja_contract(), status=_first_birja_status())
    context = admin_client.get("/birja/yuklar/").context
    assert context["default_tab"] == _first_birja_status().pk


def test_a_yuk_cannot_be_moved_into_the_other_chains_holat(admin_client):
    """Posting an Eron holat onto a birja yuk would leave the row in a bosqich its
    own list cannot draw a tab for — invisible on every view of the list it
    belongs to."""
    shipment = make_shipment(contract=_birja_contract(), status=_first_birja_status())
    chegarada = ShipmentStatus.objects.get(name="Chegarada")
    resp = admin_client.post(f"/shipments/{shipment.pk}/status/",
                             {"status": chegarada.pk})
    shipment.refresh_from_db()
    assert resp.status_code == 404
    assert shipment.status.scope == ShipmentStatus.Scope.BIRJA


def test_both_chains_end_on_the_same_arrival_holat(admin_client):
    """The one step they share. It is what turns a yuk into an ombor loti, so a
    birja load reaching it lands on the shelf exactly like an Eron one."""
    arrival = ShipmentStatus.arrival()
    assert arrival in ShipmentStatus.for_kind(birja=True)
    assert arrival in ShipmentStatus.for_kind(birja=False)

    contract = _birja_contract()
    shipment = make_shipment(contract=contract, status=_first_birja_status())
    admin_client.post(f"/shipments/{shipment.pk}/status/", {"status": arrival.pk})
    shipment.refresh_from_db()
    assert shipment.arrived is not None
    assert admin_client.get("/ombor/").status_code == 200


def test_promoting_a_holat_to_arrival_makes_it_shared(admin_client):
    """Otherwise the other chain silently loses its ending and its loads could
    never reach the ombor."""
    chegarada = ShipmentStatus.objects.get(name="Chegarada")
    admin_client.post(f"/statuses/{chegarada.pk}/edit/",
                      {"name": "Chegarada", "scope": "hamkor", "is_arrival": "on"})
    chegarada.refresh_from_db()
    assert chegarada.is_arrival and chegarada.scope == ShipmentStatus.Scope.BOTH
    assert chegarada in ShipmentStatus.for_kind(birja=True)


def test_a_new_holat_lands_at_the_end_of_its_own_chain_not_after_arrival(admin_client):
    """The arrival row sits at a deliberately high `order` so both chains end on
    it. Taking the table-wide maximum would park every new step after it."""
    admin_client.post("/statuses/new/",
                      {"name": "Birjada tekshirildi", "scope": "birja"})
    names = [s.name for s in ShipmentStatus.for_kind(birja=True)]
    assert names[-1] == "Omborga yetib keldi"
    assert names[-2] == "Birjada tekshirildi"


# --- what a birja yuk does not carry ---------------------------------------

def test_an_arrived_birja_lot_is_never_bojxonasi_tolanmagan(admin_client):
    """It crossed no border, so it owes no clearing — and since it will never carry
    a bojxona xarajat, "not recorded yet" would be permanent. The Python property
    and its SQL twin both have to say so: the row's badge comes from one and the
    pill's count from the other."""
    birja = make_shipment(contract=_birja_contract(), status=ShipmentStatus.arrival(),
                          arrived="2026-07-20")
    hamkor = make_shipment(contract=_hamkor_contract(), status=ShipmentStatus.arrival(),
                           arrived="2026-07-20")
    assert birja.customs_pending is False
    assert hamkor.customs_pending is True

    rows = admin_client.get("/shipments/", {"customs": "1"}).context["shipments"]
    assert [s.pk for s in rows] == [hamkor.pk]
    assert admin_client.get("/shipments/").context["customs_pending_count"] == 1


def test_the_birja_list_has_no_bojxona_group_and_no_qr(admin_client):
    make_shipment(contract=_birja_contract(), status=_first_birja_status())
    context = admin_client.get("/birja/yuklar/").context
    assert context["customs_pending_count"] == 0
    assert context["qr_waiting_count"] == 0
    html = admin_client.get("/birja/yuklar/").content.decode()
    assert "Bojxona to'lanmagan" not in html
    assert "QR bor" not in html


def test_the_qr_and_customs_params_narrow_nothing_on_the_birja_list(admin_client):
    """Enforced in `_filter_shipments` rather than by hiding the pills, so a
    hand-typed querystring cannot narrow a list where the answer never differs."""
    shipment = make_shipment(contract=_birja_contract(), status=_first_birja_status())
    rows = admin_client.get("/birja/yuklar/",
                            {"qr": "yoq", "customs": "1"}).context["shipments"]
    assert [s.pk for s in rows] == [shipment.pk]


def test_the_yuk_form_drops_the_qr_and_bojxonachi_boxes():
    from crm.forms import ShipmentForm
    birja_fields = ShipmentForm(birja=True).fields
    assert "qr_date" not in birja_fields and "customs_agent" not in birja_fields
    hamkor_fields = ShipmentForm(birja=False).fields
    assert "qr_date" in hamkor_fields and "customs_agent" in hamkor_fields


def test_the_yuk_detail_page_says_nothing_about_a_qr_kod(admin_client):
    """The shared detail page draws the Eron-road facts off `shipment.is_birja`.
    Left alone it printed "QR kod · Berilmagan" and offered the mark-as-given
    button — which does not read as "none was needed", it reads as a truck somebody
    forgot."""
    birja = make_shipment(contract=_birja_contract(), status=_first_birja_status())
    hamkor = make_shipment(contract=_hamkor_contract())
    assert "QR" not in admin_client.get(f"/shipments/{birja.pk}/").content.decode()
    assert "QR" in admin_client.get(f"/shipments/{hamkor.pk}/").content.decode()


def test_editing_reads_the_kind_off_the_row(admin_client):
    """A yuk never changes sides, so `shipment_edit` is shared — the form asks the
    instance which chain it is on rather than the URL."""
    from crm.forms import ShipmentForm
    shipment = make_shipment(contract=_birja_contract(), status=_first_birja_status())
    assert "qr_date" not in ShipmentForm(instance=shipment).fields
    assert admin_client.get(f"/shipments/{shipment.pk}/edit/").status_code == 200


# --- the books stay together -----------------------------------------------

def test_a_birja_lot_reaches_the_ombor_and_carries_its_xarajatlar(admin_client):
    """"Others will be together": once a birja yuk has landed it is stock like any
    other, priced by the same landed-cost rule."""
    lot = make_shipment(contract=_birja_contract(), status=ShipmentStatus.arrival(),
                        arrived="2026-07-20", kg="400").lines.first()
    ShipmentExpense.objects.create(
        shipment=lot.shipment, category=ShipmentExpense.Category.TRANSPORT,
        amount=Decimal("200"), date="2026-07-20")
    assert lot.landed_cost_per_kg > lot.contract_line.price
    html = admin_client.get("/ombor/").content.decode()
    assert "birja-1" in html


def test_a_hamkor_tolov_can_be_made_against_a_birja_kelishuv(admin_client):
    """Money owed to the birja is owed on the same page as everything else — the
    feature adds two lists, not a second set of books."""
    contract = _birja_contract(kg="1000", price="1.00")
    make_shipment(contract=contract, kg="1000")
    resp = admin_client.post("/supplier-payments/new/", supplier_payment_rows(
        {"currency": "usd", "amount": "500", "exchange_rate": "12000",
         "commission_percent": "", "method": "cash", "note": ""},
        contract=contract, date="2026-07-21"))
    assert resp.status_code == 302
    contract.refresh_from_db()
    assert contract.paid_total == Decimal("500.00")


# --- who may look ----------------------------------------------------------

@pytest.mark.parametrize("url", ["/birja/kelishuvlar/", "/birja/yuklar/",
                                 "/birja/kelishuvlar/export.xlsx",
                                 "/birja/yuklar/export.xlsx"])
def test_a_tarjimon_cannot_reach_the_birja_pages(translator_client, url):
    """Their whole job is the Eron road. `role_required` is on the shared view and
    cannot see which of the two lists it was reached through, so the guard is
    inside it."""
    assert translator_client.get(url).status_code == 403


def test_a_tarjimon_cannot_reach_a_birja_yuk_by_typing_its_number(translator_client):
    """The two per-yuk screens a tarjimon may open are shared by both pipelines, so
    without this a role that can reach no birja LIST could still reach one birja
    ROW. Their own loads are untouched."""
    birja = make_shipment(contract=_birja_contract(), status=_first_birja_status())
    hamkor = make_shipment(contract=_hamkor_contract())
    assert translator_client.get(f"/shipments/{birja.pk}/").status_code == 403
    assert translator_client.get(f"/shipments/{birja.pk}/driver/").status_code == 403
    assert translator_client.get(f"/shipments/{hamkor.pk}/").status_code == 200
    assert translator_client.get(f"/shipments/{hamkor.pk}/driver/").status_code == 200


class TestBrandAverage:
    """The o'rtacha narx under a marka a kelishuv took in several lots.

    A birja purchase buys the same granula through the day at whatever the exchange
    is asking, so one kelishuv carries и 1561 five times at five narxlar. The column
    then answers "which lot cost what" and not "what did this granula cost us",
    which is the question the operator actually has — see Contract.brand_averages.
    """

    def _lots(self, contract, brand, lots):
        for kg, price in lots:
            ContractLine.objects.create(contract=contract, brand=brand,
                                        kg=Decimal(kg), price=Decimal(price))

    def test_average_is_weighted_by_kg(self, admin_client):
        """birja-7's shape: three lots of 30 000 and a short one of 11 000. A plain
        mean over the four rows would price the 11 000 as heavily as a 30 000."""
        contract = Contract.objects.create(partner=birja_partner(),
                                           created="2026-07-01")
        self._lots(contract, "и 1561",
                   [("30000", "1.00"), ("30000", "1.20"),
                    ("30000", "1.30"), ("11000", "1.50")])
        [avg] = contract.brand_averages
        assert avg["brand"] == "и 1561"
        assert avg["kg"] == Decimal("101000.000")
        # (30 000·1 + 30 000·1.20 + 30 000·1.30 + 11 000·1.50) / 101 000 = 1.2030
        assert avg["price"] == Decimal("1.2030")
        # ...and not the plain mean over the four rows, which is 1.25.
        assert avg["price"] != Decimal("1.2500")

    def test_single_line_marka_draws_no_average(self, admin_client):
        """The ordinary kelishuv. Its average is the narx already in the column, and
        a second line repeating it under itself says nothing."""
        contract = _birja_contract(brand="LLDPE", kg="1000", price="1.00")
        assert contract.brand_averages == []

    def test_two_markalar_average_apart(self, admin_client):
        contract = Contract.objects.create(partner=birja_partner(),
                                           created="2026-07-01")
        self._lots(contract, "и 1561", [("1000", "1.00"), ("1000", "2.00")])
        self._lots(contract, "ftor oq", [("1000", "3.00"), ("3000", "5.00")])
        avgs = {g["brand"]: g["price"] for g in contract.brand_averages}
        assert avgs == {"и 1561": Decimal("1.5000"), "ftor oq": Decimal("4.5000")}

    def test_list_shows_the_average_line(self, admin_client):
        contract = Contract.objects.create(partner=birja_partner(),
                                           created="2026-07-01")
        self._lots(contract, "и 1561", [("1000", "1.00"), ("1000", "2.00")])
        html = admin_client.get("/birja/kelishuvlar/").content.decode()
        assert "o&#x27;rtacha" in html or "o'rtacha" in html

    def test_list_of_single_line_kelishuv_has_no_average_line(self, admin_client):
        _birja_contract(brand="LLDPE", kg="1000", price="1.00")
        html = admin_client.get("/birja/kelishuvlar/").content.decode()
        assert "rtacha" not in html


class TestPriceWithExpenses:
    """Xarajat bilan — the narx with everything the KELISHUV puts on a kg.

    Narx and o'rtacha answer what was agreed with the seller, which is what the
    qarz is measured in and must stay. What that kg actually cost is a second
    question, and until this column it had no answer anywhere on the kelishuv: the
    broker and the birja transport only surfaced on a yuk's tannarx, one truck at
    a time. See `Contract.expense_add_on_per_kg`.
    """

    def test_kelishuv_with_no_xarajat_prices_at_the_narx(self, admin_client):
        """Nothing has been spent, so there is nothing to add. The column repeating
        Narx is the honest answer, not a missing one."""
        contract = _birja_contract(brand="LLDPE", kg="1000", price="1.00")
        [line] = contract.lines.all()
        assert contract.expense_add_on_per_kg == Decimal("0")
        assert line.price_with_expenses == Decimal("1.0000")

    def test_broker_and_transport_both_land_on_the_kg(self, admin_client):
        """$200 of broker over 1 000 kg is 0,20/kg, and the birja transport rate is
        added whole — it is already quoted per kg."""
        contract = _birja_contract(brand="LLDPE", kg="1000", price="1.00",
                                   transport_rate_per_kg=Decimal("0.05"))
        ContractExpense.objects.create(contract=contract, amount=Decimal("200"),
                                       amount_uzs=Decimal("2400000"),
                                       exchange_rate=Decimal("12000"))
        contract = Contract.objects.get(pk=contract.pk)
        assert contract.expenses_per_kg == Decimal("0.2000")
        assert contract.expense_add_on_per_kg == Decimal("0.2500")
        [line] = contract.lines.all()
        assert line.price_with_expenses == Decimal("1.2500")

    def test_a_som_kelishuv_adds_up_on_its_own_side(self, admin_client):
        """Each side is summed from its own column — the so'm figure is not the
        dollar one re-rated, the same rule `brand_averages` follows."""
        contract = Contract.objects.create(partner=birja_partner(),
                                           created="2026-07-01",
                                           currency=Currency.UZS,
                                           transport_rate_per_kg=Decimal("350"))
        ContractLine.objects.create(contract=contract, brand="и 1561",
                                    kg=Decimal("1000"), price=Decimal("1.00"),
                                    price_uzs=Decimal("13000"))
        ContractExpense.objects.create(contract=contract, amount=Decimal("100"),
                                       amount_uzs=Decimal("1300000"),
                                       currency=Currency.UZS,
                                       exchange_rate=Decimal("13000"))
        contract = Contract.objects.get(pk=contract.pk)
        # 1 300 000 / 1 000 = 1 300 so'm/kg, plus the 350 transport.
        assert contract.expense_add_on_per_kg == Decimal("1650.00")
        [line] = contract.lines.all()
        assert line.price_with_expenses == Decimal("14650.00")

    def test_the_average_carries_the_same_add_on(self, admin_client):
        """The summary line under Xarajat bilan sits level with the one under Narx,
        so the two columns read row for row."""
        contract = Contract.objects.create(partner=birja_partner(),
                                           created="2026-07-01",
                                           transport_rate_per_kg=Decimal("0.10"))
        for price in ("1.00", "2.00"):
            ContractLine.objects.create(contract=contract, brand="и 1561",
                                        kg=Decimal("1000"), price=Decimal(price))
        [avg] = contract.brand_averages
        assert avg["price"] == Decimal("1.5000")
        assert avg["price_with_expenses"] == Decimal("1.6000")

    def test_the_list_draws_the_column(self, admin_client):
        _birja_contract(brand="LLDPE", kg="1000", price="1.00",
                        transport_rate_per_kg=Decimal("0.05"))
        html = admin_client.get("/birja/kelishuvlar/").content.decode()
        assert "Xarajat bilan" in html
        assert "1.05 $/kg" in html

    def test_a_tarjimon_sees_no_such_column(self, translator_client):
        """It is a money column like Narx beside it, and that role sees none of
        them — on the page or in the file."""
        _birja_contract(brand="LLDPE", kg="1000", price="1.00")
        resp = translator_client.get("/contracts/")
        assert resp.status_code == 200
        assert "Xarajat bilan" not in resp.content.decode()


# --- the money -----------------------------------------------------------------
#
# To'lovlar split the way the kelishuvlar and the yuklar already do. A birja to'lov
# rode the To'lovlar (hamkor) page among the Eron ones, where the single thing it
# has in common with a hamkor to'lov — that it settles a kelishuv — was the only
# thing the page said about it.

def _paid(admin_client, contract, amount="100"):
    return admin_client.post("/supplier-payments/new/", supplier_payment_rows(
        {"amount": amount}, contract=contract))


def test_the_two_tolovlar_lists_hold_only_their_own(admin_client):
    birja, hamkor = _birja_contract(), _hamkor_contract()
    _paid(admin_client, birja, "100")
    _paid(admin_client, hamkor, "200")
    on_birja = admin_client.get("/birja/tolovlar/").context["page"].object_list
    on_hamkor = admin_client.get("/supplier-payments/").context["page"].object_list
    assert [p.contract_id for p in on_birja] == [birja.pk]
    assert [p.contract_id for p in on_hamkor] == [hamkor.pk]


def test_the_totals_count_only_the_list_they_are_on(admin_client):
    """The figure under the search box is what the page is FOR — a total that
    silently included the other set of books would be the old page with a new
    title."""
    birja, hamkor = _birja_contract(), _hamkor_contract()
    _paid(admin_client, birja, "100")
    _paid(admin_client, hamkor, "200")
    [total] = admin_client.get("/birja/tolovlar/").context["totals"]
    assert total["count"] == 1 and total["paid"] == Decimal("100")


def test_the_birja_tolovlar_page_drops_the_hamkor_column_and_filter(admin_client):
    """One counterparty, so neither narrows nor says anything — the same rule the
    birja kelishuvlar list follows."""
    _paid(admin_client, _birja_contract(), "100")
    resp = admin_client.get("/birja/tolovlar/")
    assert not any(f["name"] == "partner" for f in resp.context["filters"]["fields"])
    assert "<th>Hamkor</th>" not in resp.content.decode()


def test_the_excel_button_exports_the_tolovlar_list_it_was_pressed_on(admin_client):
    """The flag is threaded through `_filter_supplier_payments` rather than read off
    the path, so the file cannot end up holding the other list. Matched on the kod —
    the workbook names the kelishuv, not its marka."""
    birja, hamkor = _birja_contract(), _hamkor_contract()
    _paid(admin_client, birja, "100")
    _paid(admin_client, hamkor, "200")
    birja_file = _sheet_text(admin_client.get("/birja/tolovlar/export.xlsx"))
    hamkor_file = _sheet_text(admin_client.get("/supplier-payments/export.xlsx"))
    assert birja.code in birja_file and hamkor.code not in birja_file
    assert hamkor.code in hamkor_file and birja.code not in hamkor_file


def test_a_tolov_goes_back_to_the_list_it_belongs_on(admin_client):
    """Create, edit and delete are one view apiece for both lists — a to'lov never
    changes sides, because its kelishuv does not — so they ask the row."""
    birja = _birja_contract()
    resp = _paid(admin_client, birja, "100")
    assert resp.status_code == 302 and resp["Location"] == "/birja/tolovlar/"
    payment = birja.supplier_payments.get()
    resp = admin_client.post(f"/supplier-payments/{payment.pk}/delete/", {})
    assert resp.status_code == 302 and resp["Location"] == "/birja/tolovlar/"


def test_the_new_tolov_modal_opens_on_the_side_it_was_opened_from(admin_client):
    """The Kelishuv picker is the modal's first question and it lists every kelishuv
    that still owes money. Opened off the birja page it must not be the wall of
    hamkor names the operator came here to get away from."""
    birja, hamkor = _birja_contract(), _hamkor_contract()
    offered = admin_client.get(
        "/supplier-payments/new/?birja=1").context["form"].fields["contract"]
    assert [c.pk for c in offered.queryset] == [birja.pk]
    # And the plain link keeps offering everything, as it always has.
    offered = admin_client.get(
        "/supplier-payments/new/").context["form"].fields["contract"]
    assert {c.pk for c in offered.queryset} == {birja.pk, hamkor.pk}


def test_a_translator_cannot_reach_the_birja_tolovlar(translator_client):
    """Their whole job is the Eron road — the same reason the birja kelishuvlar and
    yuklar lists are closed to them."""
    assert translator_client.get("/birja/tolovlar/").status_code == 403
